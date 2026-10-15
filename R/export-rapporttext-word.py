"""Exportera aktuell rapporttext utan diagram till ett gemensamt granskningsdokument."""
from pathlib import Path
from html import unescape
import hashlib
import json
import re
import subprocess
import tempfile

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
QA = Path(tempfile.gettempdir()) / 'lagesbild-segregation-kontroll' / 'word'
QA.mkdir(parents=True, exist_ok=True)
DEST = ROOT / 'granskningsunderlag'
DEST.mkdir(exist_ok=True)
OUTPUT = DEST / 'Lagesbild-segregation-rapporttext.docx'
PANDOC = Path('C:/Users/robnor2002/AppData/Local/Programs/Quarto/bin/tools/pandoc.exe')

blocks = []
placeholders = []
source_hashes = {}
headings = []


def clean_html(text):
    return unescape(re.sub(r'<[^>]+>', '', text)).strip()


def visual(match):
    html = re.sub(r'<!--.*?-->', '', match[1], flags=re.S).strip()
    if not html:
        return ''
    title = re.search(r'<h[1-6]\b[^>]*>(.*?)</h[1-6]>', html, re.S)
    if not title:
        raise ValueError('Visualisering utan identifierbar rubrik: ' + html[:100])
    title = clean_html(title[1])
    if 'id="socio-summary"' in html or 'id="lycko-cluster-table"' in html:
        label = f'[TABELL: {title}]'
    elif 'id="socio-map"' in html:
        label = f'[DIAGRAM: {title}, karta och boxplot]'
    elif 'class="knn-map-title"' in html:
        label = f'[DIAGRAM: {title}, karta]'
    else:
        label = f'[DIAGRAM: {title}]'
    placeholders.append(label)
    return '\n\n' + label + '\n\n'


def illustration(match):
    alt, path, attrs = match.groups()
    title = alt or ('Lyckoparadoxen och selektiv flyttning' if 'selektiv-flyttning' in path else Path(path).stem)
    label = f'[ILLUSTRATION: {title.rstrip(".")}]'
    placeholders.append(label)
    return label


for file in sorted((ROOT / 'kapitel').glob('*.qmd')):
    source = file.read_text(encoding='utf-8-sig')
    source_hashes[str(file.relative_to(ROOT))] = hashlib.sha256(file.read_bytes()).hexdigest()
    text = re.sub(r'\A---\s*\n.*?\n---\s*\n', '', source, count=1, flags=re.S)
    text = re.sub(r'```\{=html\}\s*\n(.*?)\n```', visual, text, flags=re.S)
    text = re.sub(r'<!--.*?-->', '', text, flags=re.S)
    text = re.sub(r'!\[([^\]]*)\]\(([^)]+)\)(\{[^\n]*\})?', illustration, text)
    # Länkar mellan kapitel pekar på samma rubrik inne i Worddokumentet.
    text = re.sub(r'\]\([^\s)]+\.qmd#([^\s)]+)\)', r'](#\1)', text)
    if file.stem == 'referenser':
        text = '# Referenser {#sec-referenser .unnumbered}\n'
    chapter = int(file.name[:2]) if file.name[:2].isdigit() else None
    subcounts = [0, 0, 0]

    def heading(match):
        marks, title, attrs = match.groups()
        level = len(marks)
        if chapter and chapter > 0:
            if level == 1:
                title = f'{chapter}. {title}'
            else:
                subcounts[level - 2] += 1
                for i in range(level - 1, len(subcounts)):
                    subcounts[i] = 0
                title = '.'.join(map(str, [chapter] + subcounts[:level - 1])) + ' ' + title
        headings.append({'level': level, 'text': title})
        return marks + ' ' + title + (attrs or '')

    text = re.sub(r'^(#{1,4})\s+(.+?)(\s+\{[^\n]*\})?\s*$', heading, text, flags=re.M)
    blocks.append(text.strip())

markdown = '\n\n'.join(blocks) + '\n'
md_file = QA / 'rapporttext.md'
md_file.write_text(markdown, encoding='utf-8')
intermediate = QA / 'rapporttext-pandoc.docx'
subprocess.run([str(PANDOC), str(md_file), '--from=markdown', '--to=docx',
                '--metadata=lang:sv-SE', '--output', str(intermediate)], check=True, cwd=ROOT)
doc = Document(intermediate)
styles = doc.styles
for style in styles:
    if style.type == 1 or style.type == 2:
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.find(qn('w:rPr'))
        if rpr is not None:
            color = rpr.find(qn('w:color'))
            if color is not None:
                for key in list(color.attrib):
                    if key != qn('w:val'):
                        del color.attrib[key]

normal = styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(8)
for name in ('Body Text', 'First Paragraph', 'Compact'):
    if name in styles:
        styles[name].font.size = Pt(11)
        styles[name].paragraph_format.line_spacing = 1.15
        styles[name].paragraph_format.space_after = Pt(8)
for level, size in [(1, 19), (2, 14), (3, 12), (4, 11)]:
    style = styles[f'Heading {level}']
    style.font.size = Pt(size)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(15)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.keep_with_next = True
    if level == 1:
        style.paragraph_format.page_break_before = True

for section in doc.sections:
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.1), Cm(2)
    section.left_margin, section.right_margin = Cm(2.3), Cm(2.3)
    section.footer_distance = Cm(1)

for paragraph in doc.paragraphs:
    paragraph.paragraph_format.widow_control = True
    if paragraph.text.startswith(('[DIAGRAM:', '[TABELL:', '[ILLUSTRATION:')):
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(12)
        paragraph.paragraph_format.keep_together = True
        for run in paragraph.runs:
            run.bold = True
            run.italic = True
            run.font.size = Pt(10)

# En kort framsida och ett TOC-fält som fylls med Word vid slutkontrollen.
first = doc.paragraphs[0]
p = first.insert_paragraph_before('Lägesbild segregation', style='Title')
p.style.font.size = Pt(28)
p.style.font.bold = True
p.paragraph_format.space_after = Pt(18)
first.insert_paragraph_before('Rapporttext för genomläsning', style='Subtitle')
first.insert_paragraph_before('Göteborgs Stad\nStadsledningskontoret')
first.insert_paragraph_before('Textversion 15 oktober 2026')
first.insert_paragraph_before('Här finns rapportens kapiteltext för kommentarer och ändringar. Diagram, kartor, tabeller och illustrationer är markerade på sina platser och kan ses i webbrapporten.')
toc_title = first.insert_paragraph_before('Innehåll')
toc_title.paragraph_format.page_break_before = True
toc_title.paragraph_format.space_after = Pt(14)
toc_title.runs[0].bold = True
toc_title.runs[0].font.size = Pt(19)
toc_p = first.insert_paragraph_before()
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'TOC \\o "1-3" \\h \\z \\u')
toc_p._p.append(field)
for n in range(1, 4):
    style_name = f'TOC {n}'
    style = styles[style_name] if style_name in styles else styles.add_style(style_name, 1)
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(10)
    style.paragraph_format.line_spacing = 1
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.left_indent = Cm((n - 1) * 0.45)

# Sidnummer förenklar hänvisningar mellan granskare och författare.
footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = 2
run = footer.add_run()
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
run._r.append(field)
doc.core_properties.title = 'Lägesbild segregation'
doc.core_properties.subject = 'Rapporttext för genomläsning'
doc.core_properties.author = 'Göteborgs Stad, Stadsledningskontoret'
doc.core_properties.language = 'sv-SE'
doc.save(OUTPUT)

# Strukturell kontroll: alla textrader och rubriker från Pandocs konvertering finns kvar.
converted = Document(intermediate)
final_texts = [p.text for p in doc.paragraphs]
offset = len(final_texts) - len(converted.paragraphs)
assert final_texts[offset:] == [p.text for p in converted.paragraphs]
assert [p.text for p in doc.paragraphs if p.style.name.startswith('Heading ')] == [h['text'] for h in headings]
assert len(placeholders) == 16, len(placeholders)
assert not doc.inline_shapes
assert not doc.tables
for file, digest in source_hashes.items():
    assert hashlib.sha256((ROOT / file).read_bytes()).hexdigest() == digest
(QA / 'exportkontroll.json').write_text(json.dumps({
    'chapters': len(blocks), 'headings': headings, 'placeholders': placeholders,
    'source_sha256': source_hashes, 'output': str(OUTPUT)
}, ensure_ascii=False, indent=2), encoding='utf-8')
print(OUTPUT)
print(f'{len(blocks)} kapitel, {len(headings)} rubriker och {len(placeholders)} bild-/tabellmarkeringar.')
